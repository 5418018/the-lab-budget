import glob
import json
import os
import subprocess
import sys
from datetime import datetime

import openpyxl
from openpyxl.utils.cell import range_boundaries


sys.stdout.reconfigure(encoding='utf-8')

base_dir = os.path.dirname(os.path.abspath(__file__))
excel_files = glob.glob(os.path.join(base_dir, '*.xlsx'))
if not excel_files:
    print('엑셀 파일을 찾을 수 없습니다. 기존 data.js를 유지합니다.')
    sys.exit(0)

def last_excel_commit(path):
    result = subprocess.run(
        ['git', 'log', '-1', '--format=%ct', '--', os.path.basename(path)],
        cwd=base_dir, capture_output=True, text=True, check=True,
    )
    return int(result.stdout.strip() or 0)


target_file = max(excel_files, key=last_excel_commit)
print(f'변환 대상 파일: {os.path.basename(target_file)}')

wb = openpyxl.load_workbook(target_file, data_only=True)

# 1. 5. 대시보드
ws_dash = wb['5. 대시보드']
dashboard_items = []
for r in range(11, 30):
    guan = ws_dash.cell(row=r, column=2).value
    item = ws_dash.cell(row=r, column=3).value
    budget = ws_dash.cell(row=r, column=4).value
    spent = ws_dash.cell(row=r, column=5).value
    rem = ws_dash.cell(row=r, column=6).value
    rate = ws_dash.cell(row=r, column=7).value
    if item:
        dashboard_items.append({
            'guan': str(guan).strip() if guan else '',
            'item': str(item).strip(),
            'budget': float(budget) if budget is not None else 0,
            'spent': float(spent) if spent is not None else 0,
            'remaining': float(rem) if rem is not None else 0,
            'rate': float(rate) if rate is not None else 0,
        })

# 2. 3. 5기 수입지출장부
ws_ledger = wb['3. 5기 수입지출장부']
ledger_table = ws_ledger.tables.get('Table_1')
if ledger_table:
    _min_col, header_row, _max_col, last_row = range_boundaries(ledger_table.ref)
    first_row = header_row + 1
else:
    first_row, last_row = 5, ws_ledger.max_row

transactions = []
for r in range(first_row, last_row + 1):
    d = ws_ledger.cell(row=r, column=1).value
    cat = ws_ledger.cell(row=r, column=2).value
    item = ws_ledger.cell(row=r, column=3).value
    inc = ws_ledger.cell(row=r, column=4).value
    exp = ws_ledger.cell(row=r, column=5).value
    bal = ws_ledger.cell(row=r, column=7).value
    desc = ws_ledger.cell(row=r, column=8).value
    pay_date = ws_ledger.cell(row=r, column=9).value
    payer = ws_ledger.cell(row=r, column=10).value
    bank = ws_ledger.cell(row=r, column=11).value

    if not any(value not in (None, '') for value in (d, cat, item, inc, exp)):
        continue

    date_str = d.strftime('%Y-%m-%d') if isinstance(d, datetime) else str(d)[:10] if d else ''
    pay_date_str = pay_date.strftime('%Y-%m-%d') if isinstance(pay_date, datetime) else str(pay_date)[:10] if pay_date else ''
    transactions.append({
        'id': len(transactions) + 1,
        'date': date_str,
        'category': str(cat).strip() if cat else '',
        'item': str(item).strip() if item else '',
        'income': float(inc) if inc is not None else 0,
        'expense': float(exp) if exp is not None else 0,
        'balance': float(bal) if bal is not None else 0,
        'description': str(desc).strip() if desc else '',
        'paymentDate': pay_date_str,
        'payer': str(payer).strip() if payer else '',
        'bank': str(bank).strip() if bank else '',
    })

# 3. 4. 5기 수입상세 및 인원변동
ws_members = wb['4. 5기 수입상세 및 인원변동']
member_stats = []
for r in range(6, 18):
    month = ws_members.cell(row=r, column=2).value
    org_count = ws_members.cell(row=r, column=3).value
    spon_count = ws_members.cell(row=r, column=4).value
    org_inc = ws_members.cell(row=r, column=5).value
    spon_inc = ws_members.cell(row=r, column=6).value
    tot_inc = ws_members.cell(row=r, column=7).value
    note = ws_members.cell(row=r, column=8).value
    if month and tot_inc is not None:
        member_stats.append({
            'month': str(month).strip(),
            'orgCount': int(org_count) if org_count is not None else 0,
            'sponsorCount': int(spon_count) if spon_count is not None else 0,
            'orgIncome': float(org_inc) if org_inc is not None else 0,
            'sponsorIncome': float(spon_inc) if spon_inc is not None else 0,
            'totalIncome': float(tot_inc),
            'note': str(note).strip() if note else '',
        })

last_balance = transactions[-1]['balance'] if transactions else 0
official_spent = sum(item['spent'] for item in dashboard_items)
total_budget = float(ws_dash['B5'].value or 0)
provisional_spent = sum(
    transaction['expense'] for transaction in transactions
    if transaction['category'] == '가예산'
)

dataset = {
    'updatedAt': datetime.now().strftime('%Y-%m-%d'),
    'title': 'THE 연구소 5기 회계 및 실시간 예산 현황',
    'kpi': {
        'totalBudget': total_budget,
        'officialSpent': official_spent,
        'provisionalSpent': provisional_spent,
        'totalSpent': official_spent + provisional_spent,
        'balance': last_balance,
        'burnRate': official_spent / total_budget if total_budget else 0,
        'fiscalElapsed': 0.67,
    },
    'expenseBudgets': dashboard_items,
    'transactions': transactions,
    'memberStats': member_stats,
}

web_dir = os.path.join(base_dir, 'web_dashboard')
os.makedirs(web_dir, exist_ok=True)

out_js = os.path.join(web_dir, 'data.js')
with open(out_js, 'w', encoding='utf-8') as f:
    f.write('window.INITIAL_BUDGET_DATA = ' + json.dumps(dataset, ensure_ascii=False, indent=2) + ';\n')

out_json = os.path.join(web_dir, 'data.json')
with open(out_json, 'w', encoding='utf-8') as f:
    json.dump(dataset, f, ensure_ascii=False, indent=2)

print(f'업데이트 완료: {out_js}, {out_json}')
